# -*- coding: utf-8 -*-
"""第 2 步增强：把 17 段角色扮演多轮对话（data/_dialogues/*.json）
展开为逐轮语料，按脱敏规则处理，输出 CSV + XLSX。

背景：用户认可用「子程序角色扮演」的方式加速生成商家与买家的经典话题对话。
这些对话是**多轮**的（8-12 条/段），比第 2 步的单轮语料更贴近真实咨询，
因此作为独立交付物补充，并沿用第 2 步的脱敏口径。

脱敏正则沿用 build_02 的经验：中文属 \\w，中文与数字间无词边界，
必须用 (?<!\\d)/(?!\\d) 而非 \\b。
"""
import glob
import json
import os
import re

import pandas as pd

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
SRC = os.path.join(DATA, "_dialogues")

# ══════════════ 脱敏 ══════════════
RE_ORDER = re.compile(r"(?<!\d)(20\d{10})(?!\d)")
RE_EXPRESS = re.compile(r"(?<!\d)(78\d{10})(?!\d)")
RE_PHONE = re.compile(r"(?<!\d)(1[3-9]\d{9})(?!\d)")
RE_IDCARD = re.compile(r"(?<!\d)(\d{17}[\dXx])(?!\d)")
RE_BANK = re.compile(r"(?<!\d)(\d{16,19})(?!\d)")
# 姓名：支持「王女士」这种无名字的「姓+称谓」，也支持「王小明先生」
RE_NAME = re.compile(
    r"(张|王|李|赵|刘|陈|杨|黄|周|吴|徐|孙|马|朱|胡|郭|林|何|高|罗|沈|韩|冯|曹|邓|许|傅|彭|曾|肖|田|董|潘|袁|蔡|蒋|余|于|杜|叶|程|魏|苏|吕|丁|任|卢|姚)"
    r"([\u4e00-\u9fa5]{0,2}?)(先生|女士|小姐|同学|老师)")
RE_WECHAT = re.compile(r"(微信|微信号|加我微|VX|vx|wechat|WeChat|QQ|qq)[：: ]*([A-Za-z0-9_\-]{5,20})")
# 地址：省市区+详细，整体抹掉只留省市
RE_ADDR_FULL = re.compile(
    r"([\u4e00-\u9fa5]{2,8}(?:省|自治区|市))([\u4e00-\u9fa5]{2,10}(?:市|区|县|州))"
    r"([\u4e00-\u9fa5]{2,20}(?:区|县|镇|乡|街道|路|街|巷|号|栋|楼|室|单元|村)[\u4e00-\u9fa5\d\-]{0,30})")


def desensitize(text):
    if not isinstance(text, str):
        return text
    # 先处理完整地址（保留省市级）
    text = RE_ADDR_FULL.sub(lambda m: m.group(1) + m.group(2) + "****", text)
    text = RE_IDCARD.sub("[已删除]", text)
    text = RE_PHONE.sub(lambda m: m.group(1)[:3] + "****" + m.group(1)[-4:], text)
    text = RE_ORDER.sub(lambda m: m.group(1)[:4] + "****" + m.group(1)[-4:], text)
    text = RE_EXPRESS.sub(lambda m: m.group(1)[:4] + "****" + m.group(1)[-4:], text)
    text = RE_BANK.sub("[已删除]", text)
    text = RE_NAME.sub(lambda m: m.group(1) + "**", text)
    text = RE_WECHAT.sub(lambda m: m.group(1) + "：[已删除]", text)
    return text


# ══════════════ 读取并展开 ══════════════
files = sorted(glob.glob(os.path.join(SRC, "d*.json")))
rows = []
summary = []

for i, fp in enumerate(files, start=1):
    with open(fp, encoding="utf-8") as f:
        d = json.load(f)
    cid = "DG{:03d}".format(i)
    turns = d.get("turns", [])
    is_handoff = bool(d.get("is_handoff", False))
    emotion = d.get("emotion", "中性")
    resolution = d.get("resolution", "已解决")
    skus = d.get("skus", [])
    n_round = len(turns) // 2

    for j, t in enumerate(turns, start=1):
        rows.append({
            "对话ID": cid,
            "话题": d.get("topic", ""),
            "意图类目": d.get("intent_category", ""),
            "渠道": d.get("channel", ""),
            "轮次": (j + 1) // 2,
            "角色": "买家" if t.get("role") == "buyer" else "客服",
            "内容": desensitize(t.get("text", "")),
            "是否转人工": "是" if is_handoff else "否",
            "用户情绪": emotion,
            "解决状态": resolution,
            "涉及SKU": "|".join(skus),
        })
    summary.append({
        "对话ID": cid,
        "话题": d.get("topic", ""),
        "意图类目": d.get("intent_category", ""),
        "渠道": d.get("channel", ""),
        "回合数": n_round,
        "发言条数": len(turns),
        "是否转人工": "是" if is_handoff else "否",
        "用户情绪": emotion,
        "解决状态": resolution,
        "涉及SKU数": len(skus),
        "涉及SKU": "|".join(skus),
    })

df = pd.DataFrame(rows)
df_sum = pd.DataFrame(summary)

df.to_csv(os.path.join(DATA, "多轮客服对话语料.csv"), index=False, encoding="utf-8-sig")

HEADERS = list(df.columns)
SUM_HEADERS = list(df_sum.columns)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
ws1 = wb.active
ws1.title = "对话语料"
ws2 = wb.create_sheet("对话汇总")

HEAD_FILL = PatternFill("solid", fgColor="C97B84")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)

# 教训：openpyxl 必须「先表头后数据」，且不跨函数分两步写
for ws, headers, data, widths, wrap_cols in [
    (ws1, HEADERS, df, [10, 22, 12, 12, 8, 8, 90, 10, 10, 10, 26], {7}),
    (ws2, SUM_HEADERS, df_sum, [10, 22, 12, 12, 8, 10, 10, 10, 10, 10, 26], {11}),
]:
    ws.append(headers)
    for rec in data.to_dict("records"):
        ws.append([rec[h] for h in headers])
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in wrap_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 26

wb.save(os.path.join(DATA, "多轮客服对话语料.xlsx"))

print("对话段数:", len(summary))
print("逐轮记录:", len(df), "条")
print("回合数合计:", df_sum["回合数"].sum())
print("转人工段数:", (df_sum["是否转人工"] == "是").sum())
print("意图类目:", df_sum["意图类目"].value_counts().to_dict())
print("渠道:", df_sum["渠道"].value_counts().to_dict())
